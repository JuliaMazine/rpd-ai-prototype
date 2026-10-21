import json
from io import BytesIO
from unittest.mock import patch

from docx import Document
from app.models import Course, CourseTeacher, Draft, Material, User


def test_saved_evidence_survives_review_export_and_edit(authenticated_client, db_session, teaching_payload, teaching_sources, monkeypatch):
    monkeypatch.setenv("LLM_PROVIDER", "ollama")
    monkeypatch.setenv("OLLAMA_MODEL", "qwen3:4b")
    user = db_session.query(User).filter_by(email="auth@test.com").one()
    course = Course(title="Курс анализа данных", description="Глубокое обучение", total_hours=40)
    db_session.add(course)
    db_session.flush()
    db_session.add(CourseTeacher(course_id=course.id, user_id=user.id))
    material = Material(course_id=course.id, filename="lecture.pdf", file_path="/missing/lecture.pdf", extracted_text="Страница 1:\n" + teaching_sources.full_text)
    db_session.add(material)
    db_session.commit()
    with patch("app.services.llm_service.ollama_service.generate", return_value=(json.dumps(teaching_payload), False)):
        response = authenticated_client.post(f"/ui/courses/{course.id}/generate-draft", follow_redirects=False)
    assert response.status_code == 303
    db_session.expire_all()
    draft = db_session.query(Draft).one()
    review = json.loads(draft.generation_metadata)
    assert review["status"] == "validated"
    assert review["evidence"][0]["location"] == "Страница 1"
    assert "глубокое обучение" in " ".join(review["warnings"])
    original = review["evidence"]
    page = authenticated_client.get(f"/drafts/{draft.id}/view")
    assert page.status_code == 200
    assert "Источники черновика" in page.text
    assert "Замечания к черновику" not in page.text
    assert review["warnings"][0] not in page.text
    assert "lecture.pdf" in page.text
    exported = authenticated_client.get(f"/drafts/{draft.id}/download-docx")
    assert exported.status_code == 200
    document = Document(BytesIO(exported.content))
    assert not any("Подтверждающие фрагменты" in p.text for p in document.paragraphs)
    assert "lecture.pdf" not in " ".join(p.text for p in document.paragraphs)
    assert len(document.tables) >= 4
    assert sum(p.text == "РАБОЧАЯ ПРОГРАММА ДИСЦИПЛИНЫ" for p in document.paragraphs) == 1
    assert not any("Статус черновика" in p.text or p.text == "Текст черновика" for p in document.paragraphs)
    assert "<table" in page.text
    response = authenticated_client.post(f"/ui/drafts/{draft.id}/edit", data={"content": "Измененный текст преподавателя"}, follow_redirects=False)
    assert response.status_code == 303
    db_session.refresh(draft)
    edited = json.loads(draft.generation_metadata)
    assert edited["edited"] is True
    assert edited["evidence"] == original
    assert "Текст изменен после генерации" in authenticated_client.get(f"/drafts/{draft.id}/view").text


def test_legacy_notebook_extraction_refreshes_cell_locations(tmp_path):
    import nbformat
    from app.services.material_service import generation_materials
    notebook = nbformat.v4.new_notebook(cells=[nbformat.v4.new_markdown_cell("Метрики качества классификации"), nbformat.v4.new_code_cell("print('never executed')")])
    path = tmp_path / "lecture.ipynb"
    path.write_text(nbformat.writes(notebook))
    material = Material(filename="lecture.ipynb", file_path=str(path), extracted_text="Legacy text without cell numbers")
    text = generation_materials([material])
    assert "Ячейка 1" in text
    assert "Ячейка 2" in text
    assert "Legacy text" not in text
