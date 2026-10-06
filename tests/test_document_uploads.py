from io import BytesIO
from pathlib import Path
from unittest.mock import patch

import pytest
import nbformat
from pptx import Presentation
from pptx.util import Inches
from docx import Document
from fastapi import HTTPException
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject

from app.models import Course, CourseTeacher, Material, User
from app.services.material_service import save_material

def pdf_bytes(*texts, password=None):
    writer = PdfWriter()
    for text in texts:
        page = writer.add_blank_page(width=600, height=800)
        if text:
            font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
            page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
            stream = DecodedStreamObject()
            stream.set_data(f"BT /F1 12 Tf 50 750 Td ({text}) Tj ET".encode())
            page[NameObject("/Contents")] = writer._add_object(stream)
    if password:
        writer.encrypt(password)
    output = BytesIO()
    writer.write(output)
    return output.getvalue()

def docx_bytes():
    document = Document()
    document.add_paragraph("Введение в дисциплину")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Тема: механика"
    table.cell(0, 1).text = "12 часов"
    document.add_paragraph("Заключение")
    output = BytesIO()
    document.save(output)
    return output.getvalue()

def notebook_bytes(code="print('hello')"):
    notebook = nbformat.v4.new_notebook(cells=[
        nbformat.v4.new_markdown_cell("# Основы Python\nПеременные и функции"),
        nbformat.v4.new_code_cell(code, outputs=[nbformat.v4.new_output("stream", name="stdout", text="Saved output should be ignored")]),
        nbformat.v4.new_raw_cell("Заключение"),
    ])
    return nbformat.writes(notebook).encode("utf-8")


def pptx_bytes():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text = "Введение в Python"
    group = slide.shapes.add_group_shape()
    group.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(1)).text = "Группированная тема"
    table = slide.shapes.add_table(1, 2, Inches(1), Inches(3), Inches(4), Inches(1)).table
    table.cell(0, 0).text = "Функции"
    table.cell(0, 1).text = "Практика"
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text = "Заключение"
    output = BytesIO()
    presentation.save(output)
    return output.getvalue()


def test_pdf_text_pages_are_extracted(db_session, sample_course):
    source = pdf_bytes("First lecture", "Second lecture")
    material = save_material(db_session, sample_course.id, "lecture.PDF", source)
    assert "First lecture" in material.extracted_text
    assert "Second lecture" in material.extracted_text
    assert material.extracted_text.index("First") < material.extracted_text.index("Second")
    assert Path(material.file_path).read_bytes() == source

def test_docx_paragraphs_and_tables_keep_document_order(db_session, sample_course):
    material = save_material(db_session, sample_course.id, "лекция.docx", docx_bytes())
    for value in ("Введение", "механика", "12 часов", "Заключение"):
        assert value in material.extracted_text
    assert material.extracted_text.index("Введение") < material.extracted_text.index("механика") < material.extracted_text.index("Заключение")

@pytest.mark.parametrize("filename,content,message", [
    ("broken.pdf", b"not a PDF", "Не удалось прочитать PDF"),
    ("broken.docx", b"not a Word document", "Не удалось прочитать документ Word"),
    ("scan.pdf", pdf_bytes(""), "В PDF не найден текст"),
    ("protected.pdf", pdf_bytes("Secret", password="secret"), "PDF защищен паролем"),
    ("old.doc", b"legacy document", "Сохраните документ Word"),
    ("image.png", b"image", "Поддерживаются файлы"),
    ("broken.pptx", b"not a presentation", "Не удалось прочитать презентацию"),
    ("broken.ipynb", b"not JSON", "Не удалось прочитать блокнот"),
    ("invalid.ipynb", b'{"cells": "invalid", "nbformat": 4, "nbformat_minor": 5}', "Не удалось прочитать блокнот"),
    ("old.ppt", b"old presentation", "Сохраните презентацию"),
], ids=["broken-pdf", "broken-docx", "scan", "encrypted", "legacy-doc", "unsupported", "broken-pptx", "broken-ipynb", "invalid-notebook", "legacy-ppt"])
def test_rejected_documents_leave_no_files_or_rows(db_session, sample_course, tmp_path, filename, content, message):
    with pytest.raises(HTTPException) as error:
        save_material(db_session, sample_course.id, filename, content)
    assert error.value.status_code == 400
    assert message in error.value.detail
    assert db_session.query(Material).count() == 0
    assert not list(tmp_path.rglob("*"))

@pytest.mark.parametrize("filename,content,expected", [
    ("lecture.pdf", pdf_bytes("Course sources"), "Course sources"),
    ("lecture.docx", docx_bytes(), "механика"),
    ("lecture.pptx", pptx_bytes(), "Группированная тема"),
    ("lecture.ipynb", notebook_bytes(), "Переменные и функции"),
], ids=["pdf", "docx", "pptx", "ipynb"])
def test_uploaded_document_text_reaches_generation(authenticated_client, db_session, filename, content, expected):
    user = db_session.query(User).filter_by(email="auth@test.com").one()
    course = Course(title="Документы курса")
    db_session.add(course)
    db_session.flush()
    db_session.add(CourseTeacher(user_id=user.id, course_id=course.id))
    db_session.commit()
    response = authenticated_client.post(f"/ui/courses/{course.id}/materials", files={"file": (filename, content)}, follow_redirects=False)
    assert response.status_code == 303
    with patch("app.routes.course.generate_draft_text", return_value="Черновик") as generate:
        response = authenticated_client.post(f"/ui/courses/{course.id}/generate-draft", follow_redirects=False)
        assert response.status_code == 303
        assert expected in generate.call_args.args[1]


def test_notebook_sources_keep_order_without_execution_or_outputs(db_session, sample_course, tmp_path):
    sentinel = tmp_path / "should-not-exist"
    code = f"from pathlib import Path; Path({str(sentinel)!r}).write_text('executed')"
    material = save_material(db_session, sample_course.id, "lecture.IPYNB", notebook_bytes(code))
    value = material.extracted_text
    assert value.index("Переменные") < value.index("Код:") < value.index("Заключение")
    assert code in value
    assert "Saved output should be ignored" not in value
    assert not sentinel.exists()


def test_presentation_groups_tables_and_slide_order(db_session, sample_course):
    material = save_material(db_session, sample_course.id, "lecture.PPTX", pptx_bytes())
    value = material.extracted_text
    assert value.index("Введение") < value.index("Группированная") < value.index("Функции") < value.index("Заключение")
    assert "Функции\tПрактика" in value
    assert "Слайд 1:" in value
    assert "Слайд 2:" in value


@pytest.mark.parametrize("extension", [".ipynb", ".pptx"])
def test_empty_notebook_and_presentation_are_rejected(db_session, sample_course, tmp_path, extension):
    if extension == ".ipynb":
        content = nbformat.writes(nbformat.v4.new_notebook()).encode()
    else:
        presentation = Presentation()
        presentation.slides.add_slide(presentation.slide_layouts[6])
        output = BytesIO()
        presentation.save(output)
        content = output.getvalue()
    with pytest.raises(HTTPException) as error:
        save_material(db_session, sample_course.id, "empty" + extension, content)
    assert error.value.status_code == 400
    assert "не найден текст" in error.value.detail
    assert db_session.query(Material).count() == 0
    assert not list(tmp_path.rglob("*"))


def test_presentation_expansion_limit_is_checked_before_saving(db_session, sample_course, tmp_path, monkeypatch):
    monkeypatch.setattr("app.services.text_extraction_service.MAX_EXPANDED_BYTES", 10)
    with pytest.raises(HTTPException) as error:
        save_material(db_session, sample_course.id, "large.pptx", pptx_bytes())
    assert "после распаковки" in error.value.detail
    assert db_session.query(Material).count() == 0
    assert not list(tmp_path.rglob("*"))
