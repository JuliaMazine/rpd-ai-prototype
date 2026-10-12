import json

import pytest

from app.services.rpd_content_service import content_prompt, parse_content, render_content, source_context

def test_each_source_gets_input_budget_including_end_of_large_file():
    materials = "[Источник: first.pdf]\n" + "A" * 24000 + "END-FIRST\n\n[Источник: second.pptx]\nMETRIC-SLIDES\n\n[Источник: third.ipynb]\nNOTEBOOK"
    context, names = source_context(materials)
    assert len(context) <= 16000
    assert "END-FIRST" in context
    assert "METRIC-SLIDES" in context
    assert "NOTEBOOK" in context
    assert names == ["first.pdf", "second.pptx", "third.ipynb"]

def test_prompt_requires_teaching_content_without_competency_codes(sample_course):
    prompt, names = content_prompt(sample_course, "Material concepts")
    assert "СФОРМУЛИРУЙ" in prompt
    assert "Коды компетенций" in prompt
    assert "Material concepts" in prompt

def test_renderer_preserves_entered_facts_and_has_plain_headings(sample_course, teaching_payload, teaching_sources):
    from decimal import Decimal
    sample_course.total_hours = 40
    sample_course.credits = Decimal("4")
    sample_course.assessment_format = "pass"
    teaching_payload["objectives"][0] = "**" + teaching_payload["objectives"][0] + "**"
    content = parse_content(json.dumps(teaching_payload), teaching_sources)
    result = render_content(sample_course, content, teaching_sources)
    assert "Общая трудоемкость\t[заполнить]" in result
    assert "Зачетные единицы: 4" in result
    assert "Зачет" in result
    assert "Экзамен" not in result
    assert "**" not in result
    assert "Соотнесение результатов обучения с компетенциями" in result
    assert "ОПК-1" not in result
    assert teaching_payload["objectives"][0].replace("**", "").lower().rstrip(".") in result.lower()
    assert "Оценить качество предсказаний на указанном наборе данных." in result
    assert "Таблица вычисленных метрик с формулами" in result
    assert "Основание:" not in result
    assert teaching_payload["topics"][0]["summary"] in result

@pytest.mark.parametrize("field", ["objectives", "knowledge", "skills", "assessment_tasks"])
def test_empty_teaching_sections_are_rejected(teaching_payload, teaching_sources, field):
    teaching_payload[field] = []
    with pytest.raises(ValueError):
        parse_content(json.dumps(teaching_payload), teaching_sources)

def test_placeholder_goals_are_rejected(teaching_payload, teaching_sources):
    teaching_payload["objectives"][0] = "уточняется преподавателем"
    with pytest.raises(ValueError):
        parse_content(json.dumps(teaching_payload), teaching_sources)

def test_unknown_source_references_are_rejected(teaching_payload, teaching_sources):
    teaching_payload["topics"][0]["claims"][0]["evidence"]["excerpt_id"] = "S3E1"
    with pytest.raises(ValueError):
        parse_content(json.dumps(teaching_payload), teaching_sources)

def test_generated_administrative_facts_are_rejected(teaching_payload, teaching_sources):
    teaching_payload["total_hours"] = 108
    with pytest.raises(ValueError):
        parse_content(json.dumps(teaching_payload), teaching_sources)


def test_reference_examples_are_separate_and_bounded(tmp_path, monkeypatch, sample_course):
    from docx import Document
    from app.services.rpd_example_service import reference_examples
    document = Document()
    document.add_paragraph("1. Цели освоения дисциплины")
    document.add_paragraph("Познакомить студентов с методами исследования операций.")
    document.add_paragraph("2. Место дисциплины")
    document.add_paragraph("Чужая образовательная программа")
    document.add_paragraph("3. Планируемые результаты")
    document.add_paragraph("ОПК-1. Чужая компетенция")
    document.add_paragraph("Уметь сравнивать алгоритмы решения оптимизационных задач.")
    document.add_paragraph("4. Трудоемкость")
    document.add_paragraph("108 часов")
    document.add_paragraph("5. Содержание дисциплины")
    document.add_paragraph("Линейное программирование и динамическое программирование.")
    document.add_paragraph("6. Литература")
    document.save(tmp_path / "РПД пример.docx")
    monkeypatch.setenv("RPD_EXAMPLES_DIR", str(tmp_path))
    example = reference_examples(str(tmp_path))
    assert "Познакомить студентов" in example
    assert "ОПК-1" not in example
    assert "108 часов" not in example
    assert len(example) <= 5500
    prompt, sources = content_prompt(sample_course, "Учебные данные новой дисциплины")
    assert "ТОЛЬКО примеры стиля" in prompt
    assert "НЕ переноси" in prompt
    assert "Познакомить студентов" in prompt
    assert sources == ["Материалы преподавателя"]


def test_model_cannot_supply_an_administrative_introduction(teaching_payload, teaching_sources):
    teaching_payload["summary"] = "Чужая программа 01.04.01 из образца РПД"
    with pytest.raises(ValueError):
        parse_content(json.dumps(teaching_payload), teaching_sources)


def test_generic_assessment_tasks_without_expected_results_are_rejected(teaching_payload, teaching_sources):
    teaching_payload["assessment_tasks"] = ["Оценить качество модели"] * 3
    with pytest.raises(ValueError):
        parse_content(json.dumps(teaching_payload), teaching_sources)


@pytest.mark.parametrize("quote", ["Пропущенные значения никогда не влияют на качество.", "Цитата из другой рабочей программы дисциплины."])
def test_fabricated_or_example_quotes_are_rejected(teaching_payload, teaching_sources, quote):
    teaching_payload["topics"][0]["claims"][0]["evidence"]["quote"] = quote
    with pytest.raises(ValueError):
        parse_content(json.dumps(teaching_payload), teaching_sources)

def test_claim_cannot_introduce_a_number(teaching_payload, teaching_sources):
    teaching_payload["topics"][0]["claims"][0]["text"] = "Заполнение пропусков улучшает качество на 25 процентов."
    with pytest.raises(ValueError):
        parse_content(json.dumps(teaching_payload), teaching_sources)

def test_source_matrix_is_preserved_without_invented_expected_metrics(sample_course, teaching_payload):
    source = "Матрица ошибок: TN=85, FP=10, FN=5, TP=9."
    sources = source_context("[Источник: metrics.txt]\n" + source)[1]
    for topic in teaching_payload["topics"]:
        topic["claims"] = [{"evidence": {"excerpt_id": "S1E1"}}]
    content = parse_content(json.dumps(teaching_payload), sources)
    result = render_content(sample_course, content, sources)
    assert source not in result
    from app.services.rpd_content_service import snapshot
    assert snapshot(sample_course, content, sources)["evidence"][0]["quote"] == source
    assert "Таблица вычисленных метрик с формулами" in result
    assert "precision =" not in result.lower()
    assert "recall =" not in result.lower()


def test_expected_numeric_answers_cannot_enter_schema(teaching_payload, teaching_sources):
    teaching_payload["assessment_tasks"][0]["expected_result"] = "Precision = 0.9; recall = 0.8"
    with pytest.raises(ValueError):
        parse_content(json.dumps(teaching_payload), teaching_sources)

def test_free_text_problem_cannot_switch_dataset(teaching_payload, teaching_sources):
    teaching_payload["assessment_tasks"][0]["problem"] = "Предскажите цены квартир по таблице продаж магазина."
    with pytest.raises(ValueError):
        parse_content(json.dumps(teaching_payload), teaching_sources)

def test_source_dataset_requires_verified_evidence(teaching_payload, teaching_sources):
    teaching_payload["assessment_tasks"][0]["dataset"]["kind"] = "source"
    with pytest.raises(ValueError):
        parse_content(json.dumps(teaching_payload), teaching_sources)

def test_evidence_snapshot_preserves_location_quote_and_hash(sample_course, teaching_payload, teaching_sources):
    from app.services.rpd_content_service import snapshot
    content = parse_content(json.dumps(teaching_payload), teaching_sources)
    review = snapshot(sample_course, content, teaching_sources)
    assert review["evidence"][0]["filename"] == "lecture.pdf"
    assert review["evidence"][0]["location"] == "Страница 1"
    assert len(review["evidence"][0]["source_sha256"]) == 64
    assert review["evidence"][0]["quote"] in review["source_excerpts"][0]["text"]

def test_missing_deep_learning_and_ambiguous_hours_are_flagged(sample_course, teaching_payload, teaching_sources):
    from app.services.rpd_content_service import coverage_warnings
    sample_course.description = "Курс включает глубокое обучение и анализ данных."
    sample_course.total_hours = 40
    sample_course.workload_confirmed = False
    content = parse_content(json.dumps(teaching_payload), teaching_sources)
    warnings = coverage_warnings(sample_course, content, teaching_sources)
    assert any("глубокое обучение" in item for item in warnings)
    assert any("контактную работу" in item for item in warnings)
    result = render_content(sample_course, content, teaching_sources)
    assert "Замечания к полноте проекта" not in result
    assert "Общая трудоемкость: 40" not in result

def test_academic_actions_are_observable():
    from app.services.rpd_content_service import academic
    assert academic("Понять метрики классификации.") == "Объяснять метрики классификации."
    assert academic("Разработать методы сокращения данных.") == "Применять методы сокращения данных."



def test_noun_based_skills_are_preserved_grammatically(teaching_payload, teaching_sources, sample_course):
    teaching_payload["skills"][0] = "Подготовка данных для обучения моделей."
    content = parse_content(json.dumps(teaching_payload), teaching_sources)
    assert "Умения: Подготовка данных" in render_content(sample_course, content, teaching_sources)

def test_knowledge_intro_accepts_punctuation():
    from app.services.rpd_content_service import knowledge_outcome
    assert knowledge_outcome("Знать, что метрики зависят от задачи.") == "Знает, что метрики зависят от задачи"


def test_raw_code_and_evidence_stay_outside_rpd(sample_course, teaching_payload):
    source = "Код:\nmail_true = np.array([0] * 100)"
    sources = source_context("[Источник: metrics.ipynb]\nЯчейка 19 (code):\n" + source)[1]
    content = parse_content(json.dumps(teaching_payload), sources)
    result = render_content(sample_course, content, sources)
    assert "mail_true" not in result
    assert "np.array" not in result
    assert "Основание:" not in result
    assert "Замечания к полноте проекта" not in result
    assert "не указана часов" not in result
    assert "№\tРаздел / тема\tСодержание\tЧасы\tВид занятий" in result

def test_every_topic_has_an_assessment(teaching_payload, teaching_sources):
    teaching_payload["assessment_tasks"] = [teaching_payload["assessment_tasks"][0]] * 2
    with pytest.raises(ValueError, match="Every topic"):
        parse_content(json.dumps(teaching_payload), teaching_sources)

def test_metrics_task_requires_true_targets_and_predictions(teaching_payload, teaching_sources):
    teaching_payload["assessment_tasks"][0]["dataset"]["kind"] = "feature_table"
    with pytest.raises(ValueError, match="true targets"):
        parse_content(json.dumps(teaching_payload), teaching_sources)

@pytest.mark.parametrize("summary", ["import numpy as np; print(data)", "Precision равен 0.9 для заданной матрицы."])
def test_syllabus_cannot_include_code_or_unverified_answers(teaching_payload, teaching_sources, summary):
    teaching_payload["topics"][0]["summary"] = summary
    with pytest.raises(ValueError):
        parse_content(json.dumps(teaching_payload), teaching_sources)

def test_reasonable_named_addition_produces_review_note_instead_of_rejection(teaching_payload, teaching_sources):
    from app.services.rpd_content_service import grounding_notes
    teaching_payload["topics"][0]["summary"] = "Вычисление MAE и MSE. Интерпретация ошибок модели."
    content = parse_content(json.dumps(teaching_payload), teaching_sources)
    notes = grounding_notes(content, teaching_sources)
    assert any("MAE" in note and "MSE" in note for note in notes)

def test_additional_topic_can_have_no_direct_evidence(teaching_payload, teaching_sources, sample_course):
    from app.services.rpd_content_service import snapshot
    teaching_payload["topics"][0]["claims"] = []
    content = parse_content(json.dumps(teaching_payload), teaching_sources)
    review = snapshot(sample_course, content, teaching_sources)
    assert any("методическое дополнение" in note for note in review["warnings"])
    assert teaching_payload["topics"][0]["summary"] in render_content(sample_course, content, teaching_sources)
