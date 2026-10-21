"""Build the shareable DOCX from report.md using the existing python-docx dependency."""
import re
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

HERE = Path(__file__).resolve().parent
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
sec.top_margin = sec.bottom_margin = Inches(0.7)
sec.left_margin = sec.right_margin = Inches(0.7)
normal = doc.styles['Normal']
normal.font.name, normal.font.size = 'Calibri', Pt(10)
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.08
for name in ['Title', 'Heading 1', 'Heading 2', 'Heading 3']:
    doc.styles[name].font.color.rgb = RGBColor.from_string('24553C')
doc.core_properties.title = 'RPD-AI: weekly progress report, 19–26 October 2026'
doc.core_properties.subject = 'Additions during 19–26 October; snapshot 21 October 2026'
doc.core_properties.author = 'RPD-AI project'

footer = sec.footer.paragraphs[0]
footer.alignment = 2
footer.add_run('RPD-AI · 19–26 October 2026     |     ')
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)

def plain(text):
    # Keep link targets visible in the portable document.
    text = re.sub(r'\[([^]]+)\]\(([^)]+)\)', r'\1 (\2)', text)
    return text.replace('**', '').replace('`', '')

def rich(p, text):
    for part in re.split(r'(\*\*.*?\*\*|`[^`]+`)', text):
        if part.startswith('**'):
            p.add_run(part[2:-2]).bold = True
        elif part.startswith('`'):
            r = p.add_run(part[1:-1])
            r.font.name, r.font.size = 'Consolas', Pt(9)
        else:
            p.add_run(plain(part))

lines = (HERE / 'report.md').read_text().splitlines()
i = 0
while i < len(lines):
    line = lines[i]
    if not line.strip():
        i += 1
        continue
    if line.startswith('```'):
        block = []
        i += 1
        while i < len(lines) and not lines[i].startswith('```'):
            block.append(lines[i])
            i += 1
        p = doc.add_paragraph()
        r = p.add_run('\n'.join(block))
        r.font.name, r.font.size = 'Consolas', Pt(8)
    elif line.startswith('|'):
        rows = []
        while i < len(lines) and lines[i].startswith('|'):
            values = [plain(s.strip()) for s in lines[i].strip('|').split('|')]
            if not all(re.fullmatch(r'[-: ]+', value) for value in values):
                rows.append(values)
            i += 1
        table = doc.add_table(rows=1, cols=len(rows[0]))
        table.style = 'Light Shading Accent 1'
        for n, values in enumerate(rows):
            cells = table.rows[0].cells if n == 0 else table.add_row().cells
            for cell, value in zip(cells, values):
                cell.text = value
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(8)
                        run.bold = n == 0
        # Repeat the header on subsequent pages; avoid splitting individual rows.
        header = OxmlElement('w:tblHeader')
        table.rows[0]._tr.get_or_add_trPr().append(header)
        for row in table.rows:
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        continue
    elif line.startswith('!['):
        match = re.fullmatch(r'!\[([^]]*)\]\(([^)]+)\)', line)
        picture_path = HERE / match[2]
        with Image.open(picture_path) as picture:
            exported_page = picture_path.name.startswith(('11-current-rpd', '12-current-rpd', '13-current-rpd'))
            max_width, max_height = (6.5, 7.35) if exported_page else (5.9, 6.15)
            width = min(max_width, max_height * picture.width / picture.height)
        doc.add_picture(str(picture_path), width=Inches(width))
        doc.paragraphs[-1].alignment = 1
    elif line.startswith('# '):
        doc.add_heading(line[2:], 0)
    elif line.startswith('## '):
        title = line[3:]
        if title.startswith('3.'):
            doc.add_page_break()
        doc.add_heading(title, 1)
    elif line.startswith('### '):
        title = line[4:]
        if title.startswith('Figure ') and not title.startswith('Figure 1 '):
            doc.add_page_break()
        doc.add_heading(title, 2)
    elif line.startswith('- '):
        rich(doc.add_paragraph(style='List Bullet'), line[2:])
    else:
        rich(doc.add_paragraph(), line)
    i += 1
doc.save(HERE / 'report.docx')
print(HERE / 'report.docx')
