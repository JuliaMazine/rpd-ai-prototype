"""Extract material text before saving the uploaded document."""
from io import BytesIO
from zipfile import ZipFile

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from fastapi import HTTPException
from pypdf import PdfReader
import nbformat
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

SUPPORTED_EXTENSIONS = {".txt", ".md", ".pdf", ".docx", ".ipynb", ".pptx"}
MAX_EXPANDED_BYTES = 100 * 1024 * 1024
MAX_PDF_PAGES = 500

def extract_material_text(extension: str, content: bytes) -> str:
    if extension == ".doc":
        raise HTTPException(400, "Сохраните документ Word в формате .docx и загрузите его снова. Старый формат .doc пока не поддерживается.")
    if extension == ".ppt":
        raise HTTPException(400, "Сохраните презентацию в формате .pptx и загрузите ее снова. Старый формат .ppt пока не поддерживается.")
    if extension not in SUPPORTED_EXTENSIONS:
        raise HTTPException(400, "Поддерживаются файлы .txt, .md, .pdf, .docx, .ipynb и .pptx.")
    if extension in {".txt", ".md"}:
        try:
            value = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            raise HTTPException(400, "Could not read file as UTF-8 text") from None
    elif extension == ".pdf":
        value = extract_pdf(content)
    elif extension == ".ipynb":
        value = extract_notebook(content)
    elif extension == ".pptx":
        value = extract_pptx(content)
    else:
        value = extract_docx(content)
    if not value.strip():
        if extension == ".pdf":
            raise HTTPException(400, "В PDF не найден текст. Если это скан, сначала распознайте его (OCR) или загрузите текстовый документ.")
        raise HTTPException(400, "В документе не найден текст. Добавьте текст или загрузите другой файл.")
    return value

def extract_pdf(content: bytes) -> str:
    try:
        reader = PdfReader(BytesIO(content))
        if reader.is_encrypted:
            raise HTTPException(400, "PDF защищен паролем. Загрузите копию без пароля.")
        if len(reader.pages) > MAX_PDF_PAGES:
            raise HTTPException(400, "PDF содержит более 500 страниц. Разделите документ на несколько файлов.")
        parts = []
        total_bytes = 0
        for page_number, page in enumerate(reader.pages, 1):
            stream = page.get_contents()
            if stream is not None:
                total_bytes += len(stream.get_data())
                if total_bytes > MAX_EXPANDED_BYTES:
                    raise HTTPException(400, "PDF слишком сложный для обработки. Разделите документ на несколько файлов.")
            page_text = page.extract_text() or ""
            if page_text.strip():
                parts.append(f"Страница {page_number}:\n{page_text}")
        return "\n\n".join(parts)
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(400, "Не удалось прочитать PDF. Проверьте файл или сохраните его заново.") from None

def block_text(container):
    for block in container.iter_inner_content():
        if isinstance(block, Paragraph):
            yield block.text
        elif isinstance(block, Table):
            for row in block.rows:
                seen = set()
                cells = []
                for cell in row.cells:
                    if cell._tc not in seen:
                        seen.add(cell._tc)
                        cells.append("\n".join(block_text(cell)))
                yield "\t".join(cells)

def extract_docx(content: bytes) -> str:
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > MAX_EXPANDED_BYTES:
                raise HTTPException(400, "Документ Word слишком большой после распаковки. Разделите его на несколько файлов.")
        document = Document(BytesIO(content))
        return "\n".join(block_text(document))
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(400, "Не удалось прочитать документ Word. Загрузите корректный файл .docx.") from None


def extract_notebook(content: bytes) -> str:
    """Read cell sources only; never execute code or include rich output payloads."""
    try:
        notebook = nbformat.reads(content.decode("utf-8-sig"), as_version=4)
        nbformat.validate(notebook)
        parts = []
        for cell_number, cell in enumerate(notebook.cells, 1):
            source = cell.source.strip()
            if not source:
                continue
            if cell.cell_type == "code":
                parts.append(f"Ячейка {cell_number} (code):\nКод:\n" + source)
            else:
                parts.append(f"Ячейка {cell_number} ({cell.cell_type}):\n" + source)
        return "\n\n".join(parts)
    except Exception:
        raise HTTPException(400, "Не удалось прочитать блокнот Jupyter. Загрузите корректный файл .ipynb в кодировке UTF-8.") from None


def slide_shape_text(shapes):
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from slide_shape_text(shape.shapes)
        elif shape.has_table:
            for row in shape.table.rows:
                yield "\t".join(cell.text for cell in row.cells if not cell.is_spanned)
        elif shape.has_text_frame:
            yield shape.text


def extract_pptx(content: bytes) -> str:
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > MAX_EXPANDED_BYTES:
                raise HTTPException(400, "Презентация слишком большая после распаковки. Разделите ее на несколько файлов.")
        presentation = Presentation(BytesIO(content))
        parts = []
        for number, slide in enumerate(presentation.slides, 1):
            body = "\n".join(value for value in slide_shape_text(slide.shapes) if value.strip())
            if body.strip():
                parts.append(f"Слайд {number}:\n{body}")
        return "\n\n".join(parts)
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(400, "Не удалось прочитать презентацию. Загрузите корректный файл .pptx.") from None
