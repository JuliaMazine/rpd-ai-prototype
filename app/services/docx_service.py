from io import BytesIO
from urllib.parse import quote

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from app.services.document_view_service import document_blocks

from app.models import Draft
from app.constants import STATUS_LABELS


def clean_filename(filename: str) -> str:
    allowed_characters = []

    for character in filename:
        if character.isalnum() or character in [" ", "_", "-", "."]:
            allowed_characters.append(character)

    cleaned = "".join(allowed_characters).strip()

    if not cleaned:
        return "draft.docx"

    if not cleaned.lower().endswith(".docx"):
        cleaned = f"{cleaned}.docx"

    return cleaned[:90]


def encode_docx_filename(filename: str) -> str:
    return quote(filename)


def add_multiline_text_to_docx(document: Document, text: str):
    lines = text.splitlines()

    for line in lines:
        stripped_line = line.strip()

        if not stripped_line:
            document.add_paragraph()
            continue

        if stripped_line.isupper() and len(stripped_line) < 90:
            paragraph = document.add_paragraph()
            run = paragraph.add_run(stripped_line)
            run.bold = True
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            continue

        if stripped_line.startswith(
            (
                "1.",
                "2.",
                "3.",
                "4.",
                "5.",
                "6.",
                "7.",
                "8.",
                "9.",
                "10.",
                "11.",
                "11.1",
                "11.2",
                "11.3",
            )
        ):
            paragraph = document.add_paragraph()
            run = paragraph.add_run(stripped_line)
            run.bold = True
            continue

        if stripped_line.startswith(("•", "-", "·")):
            document.add_paragraph(stripped_line, style="List Bullet")
            continue

        document.add_paragraph(stripped_line)


def build_draft_docx(draft: Draft) -> BytesIO:
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.left_margin, section.right_margin = Cm(3), Cm(1.5)
    section.top_margin, section.bottom_margin = Cm(2), Cm(2)
    normal = document.styles["Normal"]
    normal.font.name, normal.font.size = "Times New Roman", Pt(12)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(6)
    for name in ("Heading 1", "Heading 2"):
        document.styles[name].font.name = "Times New Roman"
        document.styles[name].font.size = Pt(12)
        document.styles[name].font.bold = True
        document.styles[name].font.color.rgb = RGBColor(0, 0, 0)
        fonts = document.styles[name].element.get_or_add_rPr().get_or_add_rFonts()
        for attribute in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
            fonts.attrib.pop(qn("w:" + attribute), None)

    full_rpd = "РАБОЧАЯ ПРОГРАММА ДИСЦИПЛИНЫ" in draft.content
    if full_rpd:
        cover = True
        for block in document_blocks(draft.content):
            if block["kind"] == "table":
                table = document.add_table(rows=1, cols=len(block["headers"]))
                table.style = "Table Grid"
                repeat_header = OxmlElement("w:tblHeader")
                table.rows[0]._tr.get_or_add_trPr().append(repeat_header)
                for cell, value in zip(table.rows[0].cells, block["headers"]):
                    cell.text = value
                    for run in cell.paragraphs[0].runs:
                        run.bold = True
                for values in block["rows"]:
                    for cell, value in zip(table.add_row().cells, values):
                        cell.text = value
                if len(block["headers"]) == 5:
                    table.autofit = False
                    for column, width in zip(table.columns, [0.8, 3.5, 7.4, 1.8, 3.0]):
                        column.width = Cm(width)
                        for cell in column.cells:
                            cell.width = Cm(width)
                document.add_paragraph()
                continue
            text = block["text"]
            kind = block["kind"]
            if text.startswith("Аннотация к"):
                cover = False
                document.add_page_break()
            if kind == "title":
                paragraph = document.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = paragraph.add_run(text)
                run.bold, run.font.size = True, Pt(14)
            elif kind in {"heading", "subheading"}:
                document.add_paragraph(text, style="Heading 1" if kind == "heading" else "Heading 2")
            elif kind == "bullet":
                document.add_paragraph(text, style="List Bullet")
            else:
                paragraph = document.add_paragraph(text)
                paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                if cover:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    if text.startswith(("Министерство", "Федеральное", "«Новосибирский", "(Новосибирский", "Факультет", "Новосибирск,")) or text == draft.course.title:
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    if text == draft.course.title:
                        paragraph.runs[0].bold = True
                        paragraph.runs[0].font.size = Pt(14)
                    if text.startswith(("Согласовано", "Декан", "«___»")):
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    else:
        # Existing hand-edited drafts retain their legacy export wrapper.
        title = document.add_paragraph("РАБОЧАЯ ПРОГРАММА ДИСЦИПЛИНЫ")
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title.runs[0].bold = True
        document.add_paragraph(draft.course.title)
        document.add_paragraph("Статус черновика: " + STATUS_LABELS.get(draft.status, draft.status))
        add_multiline_text_to_docx(document, draft.content)

    if draft.feedback:
        document.add_page_break()
        document.add_paragraph("Замечания методиста", style="Heading 1")
        add_multiline_text_to_docx(document, draft.feedback)
    output = BytesIO()
    document.save(output)
    output.seek(0)
    return output
